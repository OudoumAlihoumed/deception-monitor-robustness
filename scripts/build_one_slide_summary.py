#!/usr/bin/env python3
"""Build BASE_One_Slide_Summary.pptx — single-slide fellowship capstone summary."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "paper" / "BASE_One_Slide_Summary.pptx"
FIG = ROOT / "paper" / "figures" / "results_tension.png"

# Palette
NAVY = RGBColor(0x1A, 0x1F, 0x36)
ACCENT = RGBColor(0x2D, 0x6A, 0xDF)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x21, 0x21, 0x21)
MUTED = RGBColor(0x5C, 0x5C, 0x5C)
LIGHT_BG = RGBColor(0xF4, 0xF6, 0xFA)


def _box(slide, left, top, width, height, fill=None):
    shape = slide.shapes.add_shape(1, left, top, width, height)  # MSO_SHAPE.RECTANGLE
    shape.line.fill.background()
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    return shape


def _text(
    slide,
    left,
    top,
    width,
    height,
    text,
    *,
    size=14,
    bold=False,
    color=DARK,
    align=PP_ALIGN.LEFT,
    font_name="Calibri",
):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.name = font_name
    p.font.color.rgb = color
    return box


def _bullets(slide, left, top, width, height, heading, items, *, heading_size=13, item_size=11):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP

    p0 = tf.paragraphs[0]
    p0.text = heading
    p0.font.size = Pt(heading_size)
    p0.font.bold = True
    p0.font.name = "Calibri"
    p0.font.color.rgb = ACCENT
    p0.space_after = Pt(4)

    for item in items:
        p = tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.size = Pt(item_size)
        p.font.name = "Calibri"
        p.font.color.rgb = DARK
        p.space_after = Pt(3)
    return box


def build() -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

    # Header band
    _box(slide, Inches(0), Inches(0), Inches(13.333), Inches(1.35), NAVY)
    _text(
        slide,
        Inches(0.45),
        Inches(0.18),
        Inches(12.4),
        Inches(0.55),
        "Can We Trust Deception Monitors for AI Agents?",
        size=28,
        bold=True,
        color=WHITE,
    )
    _text(
        slide,
        Inches(0.45),
        Inches(0.72),
        Inches(12.4),
        Inches(0.35),
        "A Robustness-Gap Protocol and the Limits of Adaptive Evasion",
        size=15,
        color=RGBColor(0xC8, 0xD4, 0xE8),
    )
    _text(
        slide,
        Inches(0.45),
        Inches(1.02),
        Inches(12.4),
        Inches(0.28),
        "Oudoum A. Houmed  ·  Université Paris-Saclay  ·  Mentor: Krystal Jackson  ·  BASE Recursive Redlines Fellowship  ·  July 2026",
        size=10,
        color=RGBColor(0xA8, 0xB8, 0xD0),
    )

    # Column backgrounds
    col_w = Inches(3.05)
    gap = Inches(0.22)
    x0 = Inches(0.35)
    y = Inches(1.5)
    h = Inches(5.55)

    for i in range(3):
        _box(slide, x0 + i * (col_w + gap), y, col_w, h, LIGHT_BG)

    # Col 1 — Problem
    _bullets(
        slide,
        x0 + Inches(0.12),
        y + Inches(0.12),
        col_w - Inches(0.24),
        Inches(2.2),
        "The question",
        [
            "Deception monitors are proposed as controls against loss-of-control risk.",
            "Reported accuracies are usually measured against agents not trying to evade.",
            "I ask: how much detection survives as adversary budget rises?",
            "And when does evasion stop counting because the agent itself broke?",
        ],
        item_size=10.5,
    )
    _bullets(
        slide,
        x0 + Inches(0.12),
        y + Inches(2.45),
        col_w - Inches(0.24),
        Inches(2.9),
        "Protocol",
        [
            "Robustness gap Δdet(b) = AUROC drop vs baseline.",
            "Ladder b₀→b₄: prompt → SFT → probe-aware SFT → RL.",
            "TOST retention gate (≥80%): withhold Δdet if agent stops deceiving.",
            "Llama-3.3-70B · published Apollo probe (L22) · Modal H200 · N=60.",
        ],
        item_size=10.5,
    )

    # Col 2 — Results
    x1 = x0 + col_w + gap
    _bullets(
        slide,
        x1 + Inches(0.12),
        y + Inches(0.12),
        col_w - Inches(0.24),
        Inches(3.0),
        "What I found",
        [
            "Transfer gap (before any attack): Apollo 0.999 → agentic 0.83.",
            "Where deception is retained (b₂, b₃): probe AUROC holds (0.89, 0.85 at λ=8).",
            "Probe-aware SFT does not open a further gap within my budget.",
            "b₄ lowers AUROC (0.75) but fails retention: 0.14 (λ=8), 0.31 (λ=0.5).",
            "Softplus often attack_too_weak — reported honestly.",
        ],
        item_size=10.5,
    )
    _text(
        slide,
        x1 + Inches(0.12),
        y + Inches(3.25),
        col_w - Inches(0.24),
        Inches(0.35),
        "Primary numbers (λ=8, |C|=29)",
        size=11,
        bold=True,
        color=ACCENT,
    )
    table_shape = slide.shapes.add_table(3, 6, x1 + Inches(0.08), y + Inches(3.55), col_w - Inches(0.16), Inches(0.95))
    table = table_shape.table
    headers = ["", "b₀", "b₁", "b₂", "b₃", "b₄"]
    rows = [
        ["Probe AUROC", "0.83", "0.79", "0.89", "0.85", "0.75"],
        ["Retention", "1.00", "0.52", "0.86", "0.86", "0.14"],
    ]
    for c, htxt in enumerate(headers):
        cell = table.cell(0, c)
        cell.text = htxt
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(9)
            p.font.bold = True
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.text = val
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(9)
                if c == 5 and r == 2:
                    p.font.bold = True

    # Col 3 — Figure + conclusion
    x2 = x0 + 2 * (col_w + gap)
    if FIG.exists():
        slide.shapes.add_picture(
            str(FIG),
            x2 + Inches(0.1),
            y + Inches(0.12),
            width=col_w - Inches(0.2),
            height=Inches(2.35),
        )
    _text(
        slide,
        x2 + Inches(0.1),
        y + Inches(2.52),
        col_w - Inches(0.2),
        Inches(0.3),
        "Primary result (λ=8): detection falls only where deception collapses",
        size=9,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )
    _bullets(
        slide,
        x2 + Inches(0.12),
        y + Inches(2.9),
        col_w - Inches(0.24),
        Inches(2.5),
        "Takeaway",
        [
            "Not probe collapse — the limit of adaptive evasion under my budget.",
            "Do not import near-ceiling probe numbers from the training domain.",
            "Deliverable: protocol + Modal implementation + workshop draft (main.pdf).",
            "Acknowledgments: Krystal Jackson, BASE, GPU support via Modal.",
        ],
        item_size=10.5,
    )

    # Footer
    _text(
        slide,
        Inches(0.35),
        Inches(7.12),
        Inches(12.6),
        Inches(0.3),
        "oudoum.ali-houmed@universite-paris-saclay.fr  ·  deception-monitor-robustness",
        size=9,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    return OUT


if __name__ == "__main__":
    path = build()
    print(f"Wrote {path}")
